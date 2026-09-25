#!/usr/bin/env python3
"""
Convert a lesson plan written in plain markdown to a .docx file.

Harness-agnostic .docx export for the eng-teaching skill: any AI coding
agent that can run a shell command (Codex, Cursor, Aider, Claude Code,
etc.) can call this instead of relying on a platform-specific plugin.

Usage:
    python3 md_to_docx.py input.md output.docx

Supports:
    # / ## / ### headings         -> Title / Heading 1 / Heading 2
    **bold** and *italic* inline  -> bold/italic runs (including combined ***both***)
    - / * bullet list items       -> bulleted paragraphs
    > blockquote lines            -> distinct "note" style (italic, indented, gray)
      used for Teacher's Notes so they read as clearly separate from
      student-facing content
    | a | b | markdown tables     -> a real Word table (bold header row, gridlines) —
      used for learning-path tables
    blank lines                   -> paragraph breaks
    everything else               -> plain paragraph text

Requires: pip install python-docx
"""

import re
import sys

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
except ImportError:
    print(
        "python-docx is required. Install it with: pip install python-docx",
        file=sys.stderr,
    )
    sys.exit(1)

INLINE_PATTERN = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*)")


def add_inline_runs(paragraph, text, italic_default=False, color=None):
    """Split text on **bold**/*italic*/***bold italic*** markers and add runs."""
    pos = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            run.italic = italic_default
            if color:
                run.font.color.rgb = color

        token = match.group(0)
        if token.startswith("***") and token.endswith("***"):
            run = paragraph.add_run(token[3:-3])
            run.bold = True
            run.italic = True
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.italic = italic_default
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True

        if color:
            run.font.color.rgb = color
        pos = match.end()

    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.italic = italic_default
        if color:
            run.font.color.rgb = color


TABLE_SEPARATOR_PATTERN = re.compile(r"^\|?[\s:|-]+\|?$")


def parse_table_row(line):
    """Split a `| a | b |` row into cell strings, dropping the outer empty cells."""
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def convert(input_path, output_path):
    with open(input_path, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    note_color = RGBColor(0x55, 0x55, 0x55)  # gray, to visually separate notes

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        is_table_start = (
            stripped.startswith("|")
            and i + 1 < n
            and TABLE_SEPARATOR_PATTERN.match(lines[i + 1].strip())
        )

        if is_table_start:
            header_cells = parse_table_row(stripped)
            i += 2  # skip header + separator row
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i].strip()))
                i += 1

            table = doc.add_table(rows=1, cols=len(header_cells))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, header_cells):
                add_inline_runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True

            for row_cells in rows:
                row = table.add_row()
                for cell, text in zip(row.cells, row_cells):
                    add_inline_runs(cell.paragraphs[0], text)

            doc.add_paragraph()
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline_runs(p, stripped[2:].strip(), italic_default=True, color=note_color)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:].strip())
        elif re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, re.sub(r"^\d+\.\s", "", stripped))
        elif stripped == "---":
            doc.add_paragraph().add_run("").add_break()
        else:
            p = doc.add_paragraph()
            add_inline_runs(p, stripped)

        i += 1

    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python3 md_to_docx.py input.md output.docx", file=sys.stderr)
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
    print(f"Wrote {sys.argv[2]}")
